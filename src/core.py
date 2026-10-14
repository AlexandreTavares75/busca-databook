"""
Busca Databook - motor de leitura, indice e busca.

Le o conteudo dos documentos de uma pasta (e subpastas), guarda o texto em
um indice SQLite FTS5 e faz buscas rapidas, sem diferenciar maiusculas,
minusculas ou acentos ("trafo" encontra "TRAFO", "Trafô", "trafos"...).
"""

import hashlib
import hmac
import io
import logging
import os
import re
import shutil
import sqlite3
import sys
import threading
import time
import zipfile
from pathlib import Path

log = logging.getLogger(__name__)

# --------------------------------------------------------------------------
# Configuracao
# --------------------------------------------------------------------------

# Formatos cujo CONTEUDO e lido
EXT_CONTEUDO = {
    ".docx", ".docm", ".doc",
    ".xlsx", ".xlsm", ".xls",
    ".pptx", ".pptm",
    ".pdf",
    ".txt", ".csv", ".log", ".xml", ".json", ".md", ".htm", ".html", ".rtf",
    ".zip",
}
# Todos os demais (png, jpg, mpp, dwg, dwf, ...) sao buscados so pelo nome.

MAX_TAMANHO_MB = 300          # arquivos maiores sao indexados so pelo nome
MAX_PROFUNDIDADE_ZIP = 3      # zip dentro de zip dentro de zip...
OCR_DPI = 200
OCR_MAX_PAGINAS = 200         # limite de paginas escaneadas por PDF
SEP_ZIP = " >> "              # separador entre o zip e o arquivo interno

PREFIXO_EMPRESA = "DB_"       # pastas "DB_NOME" na pasta principal viram empresas
SENHA_MIN = 6                 # vale para senhas novas; senhas antigas continuam aceitas

FORMATO_EXPORTACAO = "buscadatabook-indice/1"
EXT_EXPORTACAO = ".bdx"

# Atributos de arquivo do Windows usados pelo OneDrive (Arquivos sob demanda)
_ATR_OFFLINE = 0x1000
_ATR_RECALL_ON_OPEN = 0x40000
_ATR_PINNED = 0x80000                  # "Sempre manter neste dispositivo"
_ATR_UNPINNED = 0x100000               # "Liberar espaco"
_ATR_RECALL_ON_DATA_ACCESS = 0x400000
_ATR_NUVEM = _ATR_OFFLINE | _ATR_RECALL_ON_OPEN | _ATR_RECALL_ON_DATA_ACCESS
_ATR_GRAVAVEIS = 0x1 | 0x2 | 0x4 | 0x20 | 0x100 | 0x2000   # somente leitura, oculto, sistema...


def pasta_dados_app() -> Path:
    """Dados do usuario (indices, usuarios, config, logs) - separados do programa."""
    base = os.environ.get("LOCALAPPDATA") or os.path.expanduser("~")
    p = Path(base) / "BuscaDatabook"
    p.mkdir(parents=True, exist_ok=True)
    return p


# --------------------------------------------------------------------------
# OCR (PDFs escaneados) - usa o Tesseract embutido no instalador
# --------------------------------------------------------------------------

_ocr_estado = {"ok": None, "idioma": "eng", "exe": None}
_ocr_lock = threading.Lock()


def pasta_tesseract_embutido() -> Path:
    if getattr(sys, "frozen", False):                       # programa instalado
        return Path(sys.executable).resolve().parent / "tesseract"
    return Path(__file__).resolve().parent.parent / "vendor" / "tesseract"  # codigo-fonte


def configurar_ocr():
    """Detecta o Tesseract (primeiro o embutido). Retorna True se o OCR estiver disponivel."""
    with _ocr_lock:
        if _ocr_estado["ok"] is not None:
            return _ocr_estado["ok"]
        try:
            import pytesseract
            embutido = pasta_tesseract_embutido()
            exe_embutido = embutido / "tesseract.exe"
            candidatos = [
                exe_embutido,
                Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe"),
                Path(r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"),
                Path(os.environ.get("LOCALAPPDATA", "")) / "Programs" / "Tesseract-OCR" / "tesseract.exe",
            ]
            for c in candidatos:
                if c.is_file():
                    pytesseract.pytesseract.tesseract_cmd = str(c)
                    if c == exe_embutido:
                        os.environ["TESSDATA_PREFIX"] = str(embutido / "tessdata")
                    _ocr_estado["exe"] = str(c)
                    break
            idiomas = pytesseract.get_languages(config="")
            _ocr_estado["idioma"] = "por+eng" if "por" in idiomas else "eng"
            _ocr_estado["ok"] = True
            log.info("OCR disponível: %s (%s)", _ocr_estado["exe"] or "PATH", _ocr_estado["idioma"])
        except Exception:
            log.warning("OCR indisponível", exc_info=True)
            _ocr_estado["ok"] = False
        return _ocr_estado["ok"]


def _ocr_imagem(img) -> str:
    import pytesseract
    return pytesseract.image_to_string(img, lang=_ocr_estado["idioma"])


# --------------------------------------------------------------------------
# Extratores de texto (recebem bytes ou caminho; devolvem str)
# --------------------------------------------------------------------------

def _abrir(fonte):
    """Aceita caminho (str/Path) ou bytes; devolve algo que as libs leem."""
    return io.BytesIO(fonte) if isinstance(fonte, (bytes, bytearray)) else fonte


def ler_docx(fonte) -> str:
    import docx
    d = docx.Document(_abrir(fonte))
    partes = [p.text for p in d.paragraphs]
    for t in d.tables:
        for linha in t.rows:
            for cel in linha.cells:
                partes.append(cel.text)
    for s in d.sections:
        for hf in (s.header, s.footer):
            try:
                partes.extend(p.text for p in hf.paragraphs)
            except Exception:
                pass
    return "\n".join(partes)


def ler_xlsx(fonte) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(_abrir(fonte), read_only=True, data_only=True)
    partes = []
    try:
        for ws in wb.worksheets:
            partes.append(f"[Aba: {ws.title}]")
            for linha in ws.iter_rows(values_only=True):
                vals = [str(v) for v in linha if v is not None and str(v).strip()]
                if vals:
                    partes.append(" | ".join(vals))
    finally:
        wb.close()
    return "\n".join(partes)


def ler_xls(fonte) -> str:
    try:
        import xlrd
    except ImportError:
        return ler_binario_legado(fonte)
    if isinstance(fonte, (bytes, bytearray)):
        wb = xlrd.open_workbook(file_contents=bytes(fonte))
    else:
        wb = xlrd.open_workbook(str(fonte))
    partes = []
    for sh in wb.sheets():
        partes.append(f"[Aba: {sh.name}]")
        for r in range(sh.nrows):
            vals = [str(v) for v in sh.row_values(r) if str(v).strip()]
            if vals:
                partes.append(" | ".join(vals))
    return "\n".join(partes)


def ler_pptx(fonte) -> str:
    from pptx import Presentation
    pres = Presentation(_abrir(fonte))
    partes = []
    for i, slide in enumerate(pres.slides, 1):
        partes.append(f"[Slide {i}]")
        for shape in slide.shapes:
            if shape.has_text_frame:
                partes.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for linha in shape.table.rows:
                    for cel in linha.cells:
                        partes.append(cel.text)
        if slide.has_notes_slide:
            partes.append(slide.notes_slide.notes_text_frame.text)
    return "\n".join(partes)


def ler_pdf(fonte, usar_ocr=True, progresso=None) -> str:
    """Le o texto do PDF. Paginas sem texto (escaneadas) passam por OCR."""
    import pypdfium2 as pdfium
    dados = bytes(fonte) if isinstance(fonte, (bytes, bytearray)) else str(fonte)
    pdf = pdfium.PdfDocument(dados)
    partes = []
    ocr_feitas = 0
    try:
        for i in range(len(pdf)):
            pag = pdf[i]
            tp = pag.get_textpage()
            texto = tp.get_text_range() or ""
            tp.close()
            if len(texto.strip()) < 20 and usar_ocr and ocr_feitas < OCR_MAX_PAGINAS and configurar_ocr():
                if progresso:
                    progresso(f"OCR página {i + 1}")
                img = pag.render(scale=OCR_DPI / 72).to_pil()
                texto = _ocr_imagem(img)
                ocr_feitas += 1
            pag.close()
            partes.append(f"[Página {i + 1}]\n{texto}")
    finally:
        pdf.close()
    return "\n".join(partes)


def ler_texto(fonte) -> str:
    dados = fonte if isinstance(fonte, (bytes, bytearray)) else Path(fonte).read_bytes()
    txt = ""
    for enc in ("utf-8", "cp1252", "latin-1"):
        try:
            txt = dados.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if txt.lstrip().startswith("{\\rtf"):
        txt = re.sub(r"\\[a-z]+-?\d* ?|[{}]", " ", txt)
    elif "<" in txt[:2000] and ">" in txt[:2000]:
        txt = re.sub(r"<[^>]+>", " ", txt)
    return txt


def ler_doc(fonte) -> str:
    """.doc antigo (Word 97-2003): tenta o Word via COM; senao, leitura bruta."""
    if not isinstance(fonte, (bytes, bytearray)) and os.name == "nt":
        try:
            import win32com.client  # pywin32 (opcional)
            word = win32com.client.DispatchEx("Word.Application")
            word.Visible = False
            try:
                d = word.Documents.Open(str(fonte), ReadOnly=True, AddToRecentFiles=False)
                txt = d.Content.Text
                d.Close(False)
                return txt
            finally:
                word.Quit()
        except Exception:
            pass
    return ler_binario_legado(fonte)


def ler_binario_legado(fonte) -> str:
    """Extrai trechos de texto legiveis de arquivos binarios antigos (.doc/.xls)."""
    dados = fonte if isinstance(fonte, (bytes, bytearray)) else Path(fonte).read_bytes()
    trechos = []
    for m in re.finditer(rb"(?:[\x20-\x7e\xc0-\xff]\x00){4,}", dados):   # UTF-16LE
        trechos.append(m.group().decode("utf-16le", "ignore"))
    for m in re.finditer(rb"[\x20-\x7e\xc0-\xff]{4,}", dados):            # cp1252
        trechos.append(m.group().decode("cp1252", "ignore"))
    return "\n".join(trechos)


LEITORES = {
    ".docx": ler_docx, ".docm": ler_docx,
    ".xlsx": ler_xlsx, ".xlsm": ler_xlsx,
    ".xls": ler_xls,
    ".doc": ler_doc,
    ".pptx": ler_pptx, ".pptm": ler_pptx,
    ".pdf": ler_pdf,
}


def extrair(nome: str, fonte, usar_ocr=True, progresso=None, profundidade=0):
    """
    Devolve lista de tuplas (subcaminho_no_zip_ou_vazio, texto).
    Para arquivos normais e uma lista com um item; para zip, um por membro.
    """
    ext = os.path.splitext(nome)[1].lower()
    if ext == ".zip":
        return _extrair_zip(fonte, usar_ocr, progresso, profundidade)
    if ext not in EXT_CONTEUDO:
        return [("", "")]
    if ext == ".pdf":
        return [("", ler_pdf(fonte, usar_ocr, progresso))]
    leitor = LEITORES.get(ext, ler_texto)
    return [("", leitor(fonte))]


def _extrair_zip(fonte, usar_ocr, progresso, profundidade):
    saida = []
    if profundidade >= MAX_PROFUNDIDADE_ZIP:
        return [("", "")]
    with zipfile.ZipFile(_abrir(fonte)) as z:
        for info in z.infolist():
            if info.is_dir():
                continue
            nome_int = _corrigir_nome_zip(info)
            ext = os.path.splitext(nome_int)[1].lower()
            texto = ""
            if ext in EXT_CONTEUDO and info.file_size <= MAX_TAMANHO_MB * 1024 * 1024:
                try:
                    dados = z.read(info)
                    sub = extrair(nome_int, dados, usar_ocr, progresso, profundidade + 1)
                    for subcam, t in sub:
                        caminho = nome_int + (SEP_ZIP + subcam if subcam else "")
                        saida.append((caminho, caminho + "\n" + t))
                    continue
                except Exception as e:
                    texto = f"(erro ao ler: {e})"
            # o nome do arquivo interno tambem e pesquisavel
            saida.append((nome_int, nome_int + "\n" + texto))
    return saida or [("", "")]


def _corrigir_nome_zip(info):
    """Zips criados no Windows costumam ter nomes em cp437/cp850."""
    nome = info.filename
    if not (info.flag_bits & 0x800):
        try:
            nome = nome.encode("cp437").decode("cp850")
        except Exception:
            pass
    return nome


# --------------------------------------------------------------------------
# OneDrive: economia de espaco
# --------------------------------------------------------------------------

def so_na_nuvem(st) -> bool:
    """True se o arquivo esta so na nuvem (OneDrive sob demanda): le-lo dispara o download."""
    atr = getattr(st, "st_file_attributes", 0)
    return bool(atr & _ATR_NUVEM) and not (atr & _ATR_PINNED)


def liberar_espaco(caminho) -> bool:
    """Devolve o arquivo a nuvem, como o "Liberar espaco" do Explorador (attrib +U -P)."""
    if os.name != "nt":
        return False
    try:
        import ctypes
        k32 = ctypes.windll.kernel32
        k32.GetFileAttributesW.argtypes = [ctypes.c_wchar_p]
        k32.GetFileAttributesW.restype = ctypes.c_uint32
        k32.SetFileAttributesW.argtypes = [ctypes.c_wchar_p, ctypes.c_uint32]
        k32.SetFileAttributesW.restype = ctypes.c_int
        atr = k32.GetFileAttributesW(str(caminho))
        if atr == 0xFFFFFFFF:
            return False
        return bool(k32.SetFileAttributesW(str(caminho), (atr & _ATR_GRAVAVEIS) | _ATR_UNPINNED))
    except Exception:
        log.debug("Não foi possível liberar espaço de %s", caminho, exc_info=True)
        return False


# --------------------------------------------------------------------------
# Empresas (pastas DB_NOME)
# --------------------------------------------------------------------------

def listar_empresas(raiz: str):
    """
    Pastas logo abaixo da raiz cujo nome comeca com "DB_".
    Devolve lista ordenada de (nome_exibido, caminho_completo). "DB_ENEL" -> "ENEL".
    """
    empresas = []
    try:
        with os.scandir(raiz) as it:
            for e in it:
                try:
                    if not e.is_dir():
                        continue
                except OSError:
                    continue
                if e.name.upper().startswith(PREFIXO_EMPRESA) and len(e.name) > len(PREFIXO_EMPRESA):
                    nome = e.name[len(PREFIXO_EMPRESA):].strip()
                    if nome:
                        empresas.append((nome, e.path))
    except OSError:
        log.warning("Não foi possível listar empresas em %s", raiz, exc_info=True)
    empresas.sort(key=lambda x: x[0].casefold())
    return empresas


# --------------------------------------------------------------------------
# Indice
# --------------------------------------------------------------------------

def caminho_indice(pasta_raiz: str) -> Path:
    h = hashlib.sha1(os.path.normcase(os.path.abspath(pasta_raiz)).encode("utf-8")).hexdigest()[:16]
    d = pasta_dados_app() / "indices"
    d.mkdir(exist_ok=True)
    return d / f"{h}.db"


ESQUEMA = """
CREATE TABLE IF NOT EXISTS arquivos (
    caminho  TEXT PRIMARY KEY,
    mtime    REAL,
    tamanho  INTEGER,
    status   TEXT
);
CREATE VIRTUAL TABLE IF NOT EXISTS conteudo USING fts5(
    caminho UNINDEXED,
    interno UNINDEXED,
    nome,
    texto,
    tokenize = "unicode61 remove_diacritics 2"
);
CREATE TABLE IF NOT EXISTS meta (chave TEXT PRIMARY KEY, valor TEXT);
"""


class Indice:
    def __init__(self, pasta_raiz: str):
        raiz = os.path.abspath(pasta_raiz)
        self.db_path = caminho_indice(raiz)
        self.con = sqlite3.connect(str(self.db_path), check_same_thread=False, timeout=30)
        self.con.execute("PRAGMA journal_mode=WAL")  # permite buscar enquanto indexa
        self.con.executescript(ESQUEMA)
        # Reaproveita a grafia da raiz gravada no indice ("C:\Dados" x "c:\dados"),
        # para que os caminhos guardados e o filtro por empresa usem o mesmo prefixo.
        r = self.con.execute("SELECT valor FROM meta WHERE chave='raiz'").fetchone()
        if r and os.path.normcase(r[0]) == os.path.normcase(raiz):
            raiz = r[0]
        else:
            self.con.execute("INSERT OR REPLACE INTO meta VALUES ('raiz', ?)", (raiz,))
            self.con.commit()
        self.raiz = raiz

    def fechar(self):
        self.con.close()

    def empresas(self):
        return listar_empresas(self.raiz)

    # ---- atualizacao incremental ----------------------------------------

    def _varrer(self, cancelado, avisar):
        """
        Lista os arquivos da pasta principal (e subpastas) sem abrir nenhum deles:
        nao dispara download do OneDrive. Devolve (no_disco, nuvem, falhas):
          no_disco: {caminho: (mtime, tamanho)}
          nuvem:    caminhos que estao so na nuvem
          falhas:   quantas pastas/arquivos nao puderam ser listados (rede, permissao...)
        """
        no_disco, nuvem, falhas = {}, set(), 0
        pilha = [self.raiz]
        while pilha:
            if cancelado():
                break
            pasta = pilha.pop()
            try:
                with os.scandir(pasta) as it:
                    entradas = list(it)
            except OSError:
                falhas += 1
                log.warning("Pasta inacessível durante a listagem: %s", pasta)
                continue
            subpastas = []
            for e in entradas:
                nome = e.name
                try:
                    if e.is_dir(follow_symlinks=False):
                        if not nome.startswith(("~", "$", ".")):   # ignora pastas ocultas/sistema
                            subpastas.append(e.path)
                        continue
                    if nome.startswith("~$"):                       # temporario do Office
                        continue
                    st = e.stat(follow_symlinks=False)
                except OSError:
                    falhas += 1
                    continue
                no_disco[e.path] = (st.st_mtime, st.st_size)
                if so_na_nuvem(st):
                    nuvem.add(e.path)
            pilha.extend(sorted(subpastas, key=str.casefold, reverse=True))  # ordem alfabetica
            avisar(f"Listando arquivos... {len(no_disco)}")
        return no_disco, nuvem, falhas

    def atualizar(self, usar_ocr=True, progresso=None, cancelado=lambda: False, liberar_nuvem=True):
        """
        Varre a pasta e (re)indexa so o que mudou.
        progresso(msg, atual, total) e chamado periodicamente.
        liberar_nuvem: arquivos que estavam so no OneDrive voltam para a nuvem depois de lidos.
        Devolve dict com contagens.
        """
        def avisar(msg, i=0, n=0):
            if progresso:
                progresso(msg, i, n)

        avisar("Listando arquivos...")
        no_disco, nuvem, falhas = self._varrer(cancelado, avisar)

        if cancelado():
            # varredura incompleta: nao remove nada do indice com base nela
            return {"total": len(no_disco), "lidos": 0, "removidos": 0, "erros": 0,
                    "liberados": 0, "falhas": falhas, "cancelado": True}

        no_indice = {r[0]: (r[1], r[2]) for r in
                     self.con.execute("SELECT caminho, mtime, tamanho FROM arquivos")}

        # Se alguma pasta nao pode ser lida (rede/VPN instavel), nao remove nada:
        # um arquivo "sumido" pode estar apenas inacessivel naquele momento.
        removidos = [] if falhas else [p for p in no_indice if p not in no_disco]
        novos = [p for p, mt in no_disco.items() if _mudou(no_indice.get(p), mt)]

        for p in removidos:
            self._remover(p)
        self.con.commit()

        erros = liberados = 0
        total = len(novos)
        t_commit = time.time()
        for i, p in enumerate(novos, 1):
            if cancelado():
                break
            nome = os.path.basename(p)
            avisar(f"Lendo {nome}", i, total)
            mtime, tam = no_disco[p]
            status = "ok"
            leu_conteudo = (tam <= MAX_TAMANHO_MB * 1024 * 1024
                            and os.path.splitext(nome)[1].lower() in EXT_CONTEUDO)
            try:
                if tam > MAX_TAMANHO_MB * 1024 * 1024:
                    partes = [("", "")]
                    status = "grande"
                else:
                    partes = extrair(nome, p, usar_ocr,
                                     lambda m, _n=nome, _i=i: avisar(f"{_n}: {m}", _i, total))
            except Exception as e:
                partes = [("", "")]
                status = f"erro: {type(e).__name__}: {e}"[:300]
                erros += 1
            # so devolve a nuvem o que estava na nuvem antes e foi baixado para leitura
            if liberar_nuvem and leu_conteudo and p in nuvem and liberar_espaco(p):
                liberados += 1
            self._remover(p)
            nome_busca = _nome_pesquisavel(p, self.raiz)
            for interno, texto in partes:
                self.con.execute(
                    "INSERT INTO conteudo (caminho, interno, nome, texto) VALUES (?,?,?,?)",
                    (p, interno, nome_busca, texto))
            self.con.execute("INSERT OR REPLACE INTO arquivos VALUES (?,?,?,?)",
                             (p, mtime, tam, status))
            if time.time() - t_commit > 5:
                self.con.commit()
                t_commit = time.time()
        self.con.commit()
        self.con.execute("INSERT OR REPLACE INTO meta VALUES ('atualizado', ?)",
                         (time.strftime("%d/%m/%Y %H:%M"),))
        self.con.commit()
        return {"total": len(no_disco), "lidos": total, "removidos": len(removidos),
                "erros": erros, "liberados": liberados, "falhas": falhas, "cancelado": cancelado()}

    def contem(self, caminho) -> bool:
        """True se o caminho esta dentro da pasta principal deste indice."""
        raiz = os.path.normcase(os.path.join(self.raiz, ""))
        return os.path.normcase(os.path.abspath(caminho)).startswith(raiz)

    def _remover(self, p):
        self.con.execute("DELETE FROM conteudo WHERE caminho = ?", (p,))
        self.con.execute("DELETE FROM arquivos WHERE caminho = ?", (p,))

    def info(self):
        n = self.con.execute("SELECT COUNT(*) FROM arquivos").fetchone()[0]
        r = self.con.execute("SELECT valor FROM meta WHERE chave='atualizado'").fetchone()
        return n, (r[0] if r else None)

    def arquivos_com_erro(self):
        return self.con.execute(
            "SELECT caminho, status FROM arquivos WHERE status LIKE 'erro%' ORDER BY caminho").fetchall()

    # ---- busca ------------------------------------------------------------

    def _prefixo_seguro(self, pasta):
        """Prefixo da pasta (com barra final), garantindo que ela esta dentro da raiz."""
        prefixo = os.path.join(os.path.abspath(pasta), "")
        raiz = os.path.join(self.raiz, "")
        if not os.path.normcase(prefixo).startswith(os.path.normcase(raiz)):
            raise ValueError("A pasta escolhida está fora da pasta principal.")
        # usa a grafia da raiz gravada + o restante do caminho
        return raiz + prefixo[len(raiz):]

    def buscar(self, termo: str, pasta: str = None, limite=2000):
        """
        Regras:
          trafo            -> palavras que comecam com "trafo" (trafo, trafos...)
          "trafo 500 kva"  -> frase exata
          trafo bucha      -> documentos com as duas palavras
          trafo OU bucha   -> qualquer uma das palavras
        pasta: se informada, restringe a busca a essa pasta e subpastas (ex.: uma empresa).
        Devolve lista de dicts: caminho, interno, nome, pasta, ext, trecho, onde, mtime, tamanho.
        """
        consulta = montar_consulta(termo)
        if not consulta:
            return []
        filtro, params = "", [consulta]
        if pasta:
            prefixo = self._prefixo_seguro(pasta)
            # comparacao exata de prefixo (LIKE trataria "_" de "DB_" como curinga)
            filtro = " AND substr(caminho, 1, ?) = ?"
            params += [len(prefixo), prefixo]
        params.append(limite)
        sql = f"""
            SELECT caminho, interno,
                   snippet(conteudo, 3, '[', ']', ' … ', 14) AS trecho,
                   snippet(conteudo, 2, '[', ']', ' … ', 14) AS trecho_nome,
                   bm25(conteudo, 0, 0, 5.0, 1.0) AS rank
            FROM conteudo
            WHERE conteudo MATCH ?{filtro}
            ORDER BY rank
            LIMIT ?
        """
        resultados = []
        vistos = set()
        for caminho, interno, trecho, trecho_nome, _ in self.con.execute(sql, params):
            chave = (caminho, interno)
            if chave in vistos:
                continue
            vistos.add(chave)
            achou_no_texto = "[" in (trecho or "")
            resultados.append({
                "caminho": caminho,
                "interno": interno,
                "nome": os.path.basename(caminho),
                "pasta": os.path.dirname(caminho),
                "ext": os.path.splitext(interno or caminho)[1].lower().lstrip("."),
                "trecho": _limpar(trecho if achou_no_texto else trecho_nome),
                "onde": "conteúdo" if achou_no_texto else "nome",
            })
        meta = self._meta_arquivos(r["caminho"] for r in resultados)
        for r in resultados:
            r["mtime"], r["tamanho"] = meta.get(r["caminho"], (0.0, 0))
        return resultados

    def _meta_arquivos(self, caminhos):
        """{caminho: (mtime, tamanho)} - consulta em lotes (limite de parametros do SQLite)."""
        meta, lista = {}, list(set(caminhos))
        for i in range(0, len(lista), 500):
            parte = lista[i:i + 500]
            marcas = ",".join("?" * len(parte))
            for c, mt, tam in self.con.execute(
                    f"SELECT caminho, mtime, tamanho FROM arquivos WHERE caminho IN ({marcas})", parte):
                meta[c] = (mt or 0.0, tam or 0)
        return meta

    # ---- exportar / importar ----------------------------------------------

    def _filtro_exportacao(self, pastas, incluir_fora):
        """WHERE para exportar so algumas empresas. pastas=None exporta o indice inteiro."""
        if pastas is None:
            return "1", []
        conds, params = [], []
        for p in pastas:
            pref = self._prefixo_seguro(p)
            conds.append("substr(caminho, 1, ?) = ?")
            params += [len(pref), pref]
        if incluir_fora:   # arquivos que nao estao em nenhuma pasta DB_
            todas = [self._prefixo_seguro(c) for _, c in self.empresas()]
            if todas:
                conds.append("(" + " AND ".join("substr(caminho, 1, ?) <> ?" for _ in todas) + ")")
                for t in todas:
                    params += [len(t), t]
            else:
                conds.append("1")
        if not conds:
            raise ValueError("Escolha ao menos uma empresa para exportar.")
        return "(" + " OR ".join(conds) + ")", params

    def exportar(self, destino, pastas=None, incluir_fora=True):
        """
        Gera um arquivo .bdx (SQLite) com o indice inteiro ou so das empresas escolhidas.
        Devolve quantos arquivos foram exportados.
        """
        destino = Path(destino)
        tmp = destino.with_name(destino.name + ".tmp")
        if tmp.exists():
            tmp.unlink()
        filtro, params = self._filtro_exportacao(pastas, incluir_fora)
        con = sqlite3.connect(str(tmp))
        try:
            con.executescript(ESQUEMA)
            con.execute("ATTACH DATABASE ? AS origem", (str(self.db_path),))
            con.execute(f"INSERT INTO arquivos SELECT caminho, mtime, tamanho, status "
                        f"FROM origem.arquivos WHERE {filtro}", params)
            con.execute(f"INSERT INTO conteudo (caminho, interno, nome, texto) "
                        f"SELECT caminho, interno, nome, texto FROM origem.conteudo WHERE {filtro}", params)
            r = con.execute("SELECT valor FROM origem.meta WHERE chave='atualizado'").fetchall()
            con.executemany("INSERT OR REPLACE INTO meta VALUES (?, ?)", [
                ("raiz", self.raiz),
                ("formato", FORMATO_EXPORTACAO),
                ("exportado_em", time.strftime("%d/%m/%Y %H:%M")),
                ("atualizado", r[0][0] if r else ""),
            ])
            con.commit()
            n = con.execute("SELECT COUNT(*) FROM arquivos").fetchall()[0][0]
            con.execute("DETACH DATABASE origem")
        except Exception:
            con.close()
            tmp.unlink(missing_ok=True)
            raise
        con.close()
        os.replace(tmp, destino)
        log.info("Índice exportado: %s (%d arquivos)", destino, n)
        return n

    def mesclar(self, arquivo):
        """Junta um indice (ja com os caminhos desta pasta) ao indice atual.
        Arquivos que ja existiam sao substituidos; os demais sao mantidos."""
        self.con.commit()
        self.con.execute("ATTACH DATABASE ? AS imp", (str(arquivo),))
        try:
            with self.con:
                self.con.execute("DELETE FROM conteudo WHERE caminho IN (SELECT caminho FROM imp.arquivos)")
                self.con.execute("DELETE FROM arquivos WHERE caminho IN (SELECT caminho FROM imp.arquivos)")
                self.con.execute("INSERT INTO arquivos SELECT caminho, mtime, tamanho, status FROM imp.arquivos")
                self.con.execute("INSERT INTO conteudo (caminho, interno, nome, texto) "
                                 "SELECT caminho, interno, nome, texto FROM imp.conteudo")
                r = self.con.execute("SELECT valor FROM imp.meta WHERE chave='atualizado'").fetchall()
                if r and r[0][0]:
                    self.con.execute("INSERT OR REPLACE INTO meta VALUES ('atualizado', ?)", (r[0][0],))
            n = self.con.execute("SELECT COUNT(*) FROM imp.arquivos").fetchall()[0][0]
        finally:
            self.con.execute("DETACH DATABASE imp")
        return n


def ler_exportacao(arquivo):
    """Le e valida o cabecalho de um indice exportado. Devolve dict(raiz, arquivos, exportado_em)."""
    try:
        con = sqlite3.connect(str(arquivo))
    except sqlite3.Error as e:
        raise ValueError(f"Não foi possível abrir o arquivo: {e}")
    try:
        con.execute("PRAGMA query_only = 1")
        meta = dict(con.execute("SELECT chave, valor FROM meta").fetchall())
        if meta.get("formato") != FORMATO_EXPORTACAO:
            raise ValueError("Este arquivo não é um índice exportado pelo Busca Databook.")
        n = con.execute("SELECT COUNT(*) FROM arquivos").fetchone()[0]
    except sqlite3.DatabaseError as e:
        raise ValueError(f"Arquivo de índice inválido ou corrompido: {e}")
    finally:
        con.close()
    return {"raiz": meta.get("raiz", ""), "arquivos": n, "exportado_em": meta.get("exportado_em", "")}


def importar_indice(arquivo, nova_raiz):
    """
    Importa um .bdx para a pasta principal nova_raiz (pode ser outro caminho que o de origem:
    os caminhos sao ajustados). Junta ao indice existente dessa pasta. Devolve quantos arquivos.
    """
    info = ler_exportacao(arquivo)
    idx = Indice(nova_raiz)
    tmp = pasta_dados_app() / "indices" / f"importando-{os.getpid()}-{int(time.time())}.db"
    try:
        shutil.copyfile(arquivo, tmp)   # trabalha numa copia: o arquivo recebido fica intacto
        con = sqlite3.connect(str(tmp))
        try:
            if con.execute("PRAGMA quick_check").fetchone()[0] != "ok":
                raise ValueError("O arquivo de índice está corrompido.")
            p_ant = os.path.join(info["raiz"], "")
            p_nova = os.path.join(idx.raiz, "")
            if p_ant != p_nova:
                for tabela in ("arquivos", "conteudo"):
                    con.execute(f"UPDATE {tabela} SET caminho = ? || substr(caminho, ?) "
                                f"WHERE substr(caminho, 1, ?) = ?",
                                (p_nova, len(p_ant) + 1, len(p_ant), p_ant))
            # seguranca: nenhum caminho pode apontar para fora da pasta escolhida
            fora = sum(con.execute(f"SELECT COUNT(*) FROM {t} WHERE substr(caminho, 1, ?) <> ?",
                                   (len(p_nova), p_nova)).fetchone()[0] for t in ("arquivos", "conteudo"))
            if fora:
                raise ValueError("O índice contém caminhos fora da pasta escolhida. Importação cancelada.")
            con.commit()
        finally:
            con.close()
        n = idx.mesclar(tmp)
        log.info("Índice importado de %s para %s (%d arquivos)", arquivo, idx.raiz, n)
        return n
    finally:
        idx.fechar()
        try:
            tmp.unlink()
        except OSError:
            pass


# --------------------------------------------------------------------------
# Copias e versoes repetidas nos resultados
# --------------------------------------------------------------------------

def _mudou(no_indice, no_disco) -> bool:
    """Compara (mtime, tamanho) com tolerancia de 1 ms na data."""
    if no_indice is None:
        return True
    return no_indice[1] != no_disco[1] or abs((no_indice[0] or 0) - no_disco[0]) > 0.001


def _nome_documento(r) -> str:
    """Nome do documento (para zip: nome do arquivo interno)."""
    if r["interno"]:
        ultimo = r["interno"].split(SEP_ZIP)[-1]
        return ultimo.replace("\\", "/").rsplit("/", 1)[-1]
    return r["nome"]


def _empresa_do_caminho(caminho, prefixo_raiz) -> str:
    """Pasta DB_ de primeiro nivel onde o arquivo esta ("" se nao estiver em nenhuma)."""
    if not os.path.normcase(caminho).startswith(os.path.normcase(prefixo_raiz)):
        return ""
    primeira = re.split(r"[\\/]", caminho[len(prefixo_raiz):], maxsplit=1)[0]
    return primeira.casefold() if primeira.upper().startswith(PREFIXO_EMPRESA) else ""


def agrupar_repetidos(resultados, raiz):
    """
    Junta documentos com o mesmo nome dentro da mesma empresa e mostra so um:
      - versoes (tamanhos diferentes): fica a mais recente;
      - copias identicas (mesmo tamanho): fica a do caminho mais curto.
    Cada item devolvido traz em "copias" todos os documentos do grupo (mais recente primeiro).
    A ordem de relevancia da busca e mantida.
    """
    prefixo = os.path.join(raiz, "")
    grupos, ordem = {}, []
    for r in resultados:
        chave = (_empresa_do_caminho(r["caminho"], prefixo), _nome_documento(r).casefold())
        if chave not in grupos:
            grupos[chave] = []
            ordem.append(chave)
        grupos[chave].append(r)
    saida = []
    for chave in ordem:
        membros = grupos[chave]
        versoes = {}
        for r in membros:
            versoes.setdefault(r.get("tamanho", 0), []).append(r)
        recente = max(versoes.values(), key=lambda v: max(x.get("mtime", 0) for x in v))
        principal = dict(min(recente, key=lambda x: (len(x["caminho"]) + len(x["interno"]), x["caminho"])))
        principal["copias"] = sorted(membros, key=lambda x: (-x.get("mtime", 0), len(x["caminho"])))
        saida.append(principal)
    return saida


def _nome_pesquisavel(caminho, raiz):
    """Nome do arquivo + subpastas relativas, para achar tambem por pasta."""
    rel = os.path.relpath(caminho, raiz)
    base = os.path.splitext(rel)[0]
    # separa palavras grudadas por _ - . para o tokenizador
    return re.sub(r"[_\-.\\/]+", " ", base)


def _limpar(t):
    return re.sub(r"\s+", " ", t or "").strip()


def montar_consulta(termo: str) -> str:
    """Converte o texto digitado em consulta FTS5 segura (sem operadores do usuario)."""
    termo = (termo or "").strip()
    if not termo:
        return ""
    partes = re.findall(r'"[^"]+"|\S+', termo)
    grupos, atual = [], []
    for p in partes:
        if p.upper() in ("OU", "OR"):
            if atual:
                grupos.append(atual)
            atual = []
            continue
        if p.startswith('"') and p.endswith('"') and len(p) > 2:
            frase = " ".join(re.findall(r"\w+", p[1:-1], flags=re.UNICODE))
            if frase:
                atual.append('"' + frase + '"')
        else:
            for tok in re.findall(r"\w+", p, flags=re.UNICODE):
                atual.append('"' + tok + '"*')
    if atual:
        grupos.append(atual)
    grupos = [g for g in grupos if g]
    if not grupos:
        return ""
    return " OR ".join("(" + " AND ".join(g) + ")" for g in grupos)


# --------------------------------------------------------------------------
# Usuarios (login)
# --------------------------------------------------------------------------

class Usuarios:
    ITERACOES = 200_000

    def __init__(self):
        self.con = sqlite3.connect(str(pasta_dados_app() / "usuarios.db"))
        self.con.execute("""CREATE TABLE IF NOT EXISTS usuarios (
            login TEXT PRIMARY KEY, sal BLOB, hash BLOB, admin INTEGER)""")
        self.con.commit()

    @classmethod
    def _hash(cls, senha, sal):
        return hashlib.pbkdf2_hmac("sha256", senha.encode("utf-8"), sal, cls.ITERACOES)

    @staticmethod
    def _checar_senha(senha):
        if len(senha) < SENHA_MIN:
            raise ValueError(f"A senha precisa ter pelo menos {SENHA_MIN} caracteres.")

    def vazio(self):
        return self.con.execute("SELECT COUNT(*) FROM usuarios").fetchone()[0] == 0

    def criar(self, login, senha, admin=False):
        login = login.strip().lower()
        if not login:
            raise ValueError("Informe um nome de usuário.")
        self._checar_senha(senha)
        if self.con.execute("SELECT 1 FROM usuarios WHERE login=?", (login,)).fetchone():
            raise ValueError("Já existe um usuário com esse nome.")
        sal = os.urandom(16)
        self.con.execute("INSERT INTO usuarios VALUES (?,?,?,?)",
                         (login, sal, self._hash(senha, sal), int(admin)))
        self.con.commit()

    def validar(self, login, senha):
        login = login.strip().lower()
        r = self.con.execute("SELECT sal, hash, admin FROM usuarios WHERE login=?", (login,)).fetchone()
        if r:
            if hmac.compare_digest(self._hash(senha, r[0]), r[1]):
                return {"login": login, "admin": bool(r[2])}
        else:
            # mesmo custo de tempo para usuario inexistente (nao revela quais logins existem)
            self._hash(senha, b"\0" * 16)
        return None

    def listar(self):
        return self.con.execute("SELECT login, admin FROM usuarios ORDER BY login").fetchall()

    def remover(self, login):
        self.con.execute("DELETE FROM usuarios WHERE login=?", (login,))
        self.con.commit()

    def trocar_senha(self, login, nova):
        self._checar_senha(nova)
        sal = os.urandom(16)
        self.con.execute("UPDATE usuarios SET sal=?, hash=? WHERE login=?",
                         (sal, self._hash(nova, sal), login))
        self.con.commit()